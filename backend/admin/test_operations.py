import hashlib
import json
import os
import stat
import sqlite3
import tempfile
import unittest
import asyncio
from datetime import datetime
from pathlib import Path
from unittest.mock import patch
from contextlib import nullcontext
from types import SimpleNamespace
from docx import Document
from fastapi import HTTPException

from backend.admin import operations as ops
from backend.api.routes import operations as operations_route


class OperationsTests(unittest.TestCase):
    def setUp(self):
        ops._worker_stop.clear()
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.db = self.root / "operations.sqlite"
        self.docs = self.root / "documents"
        self.docs.mkdir()

    def tearDown(self):
        ops._worker_stop.set()
        ops._worker_wake.clear()
        self.tmp.cleanup()

    def test_vk_source_validation_and_manual_job_history(self):
        source = ops.add_vk_source("https://vk.ru/Dainima", self.db)
        self.assertEqual(source["id"], "vk:dainima")
        self.assertEqual(source["url"], "https://vk.com/dainima")
        with self.assertRaises(ValueError): ops.add_vk_source("https://evil.test/dainima", self.db)
        with self.assertRaises(ValueError): ops.add_vk_source("https://vk.com/spoyunost", self.db)
        with self.assertRaisesRegex(ValueError, "через форму ссылки на пост"):
            ops.add_vk_source("https://vk.com/wall-55_101", self.db)
        job = ops.create_vk_job(source["id"], 12, self.db)
        self.assertEqual(job["status"], "queued")
        self.assertEqual(job["result"], {"limit": 12})
        con = ops.connect(self.db); con.execute("UPDATE operations_jobs SET status='failed' WHERE id=?", (job["id"],)); con.commit(); con.close()
        retry = ops.retry_job(job["id"], self.db)
        self.assertEqual(retry["status"], "queued")
        self.assertEqual(retry["retry_of"], job["id"])
        with self.assertRaises(ValueError): ops.retry_job(retry["id"], self.db)

    def test_existing_group_sources_import_offline_idempotently_with_titles_and_raw_baseline(self):
        links = self.root / "group_links.txt"
        links.write_text("#https://vk.ru/commented\nhttps://vk.ru/dainima\nhttps://t.me/example\nhttps://vk.ru/kgeu_official\nhttps://vk.ru/rso_rt\n", encoding="utf-8")
        groups = self.root / "groups"
        groups.mkdir()
        (groups / "groups_dainima.json").write_text('{"name":"ССО Дайнима"}', encoding="utf-8")
        posts = self.root / "posts"
        posts.mkdir()
        (posts / "posts_dainima.jsonl").write_text('{"published_at":"2026-10-16T10:00:00+00:00"}\n{"published_at":"2026-10-17T12:00:00+00:00"}\n', encoding="utf-8")
        with patch("backend.admin.operations.STORAGE_DIR", self.root):
            first = ops.import_existing_vk_sources(links, self.db)
            second = ops.import_existing_vk_sources(links, self.db)
        self.assertEqual(first["imported"], ["vk:dainima"])
        self.assertEqual(len(first["skipped"]), 2)
        self.assertTrue(all(item["reason"] == "excluded" for item in first["skipped"]))
        self.assertEqual(second["already_present"], ["vk:dainima"])
        self.assertEqual(ops.list_jobs(20, self.db), [])
        source = next(item for item in ops.list_sources(self.db) if item["id"] == "vk:dainima")
        self.assertEqual(source["title"], "ССО Дайнима")
        self.assertEqual(source["last_collected_at"], "2026-10-17T12:00:00+00:00")

    def test_individual_vk_permalink_creates_one_source_only_job_and_inbox_item(self):
        source, job = ops.create_vk_post_source("https://vk.ru/wall-55_101", self.db)
        self.assertEqual(source["id"], "vkpost:wall-55_101")
        self.assertEqual(job["kind"], "vk_post_fetch")
        with self.assertRaises(ValueError): ops.create_vk_post_source("https://evil.test/wall-55_101", self.db)
        with self.assertRaises(ValueError): ops.create_vk_post_source("https://vk.com/dainima", self.db)

        class FakeWall:
            def getById(self, posts):
                self.posts = posts
                return [{"id":101,"owner_id":-55,"date":1792218000,"text":"Текст поста",
                         "attachments":[]}]
        class FakeExtractor:
            def __init__(self, token): self.api = SimpleNamespace(wall=FakeWall())
        with patch("backend.admin.secrets.get_secret", return_value="test-token"):
            result = ops._collect_vk_post({**source,"kind":"vk_post"}, job["id"], self.db,
                                          FakeExtractor, lambda text: (text, [], []))
        self.assertEqual(result["fetched"], 1)
        item = ops.list_inbox("queued", 5, self.db)[0]
        self.assertEqual(item["source_kind"], "vk_post")
        self.assertEqual(item["source_url"], "https://vk.com/wall-55_101")
        self.assertEqual(ops.get_inbox_item(item["id"], self.db)["text"], "Текст поста")

    def test_docx_preserves_paragraph_table_order_and_pdf_keeps_page_markers(self):
        doc_id = "12345678-1234-5678-1234-567812345682"
        path = self.docs / f"{doc_id}.docx"
        document = Document()
        document.add_paragraph("BEFORE_TABLE")
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "TABLE_CONTENT"
        document.add_paragraph("AFTER_TABLE")
        document.save(path)
        docx_text, _ = ops._extract_document(path, self.docs)
        self.assertLess(docx_text.index("BEFORE_TABLE"), docx_text.index("TABLE_CONTENT"))
        self.assertLess(docx_text.index("TABLE_CONTENT"), docx_text.index("AFTER_TABLE"))

        pdf_id = "12345678-1234-5678-1234-567812345683"
        pdf_path = self.docs / f"{pdf_id}.pdf"
        pdf_path.write_bytes(b"fake-pdf")
        class FakePdf:
            pages = [SimpleNamespace(extract_text=lambda layout=False: "PAGE_A"),
                     SimpleNamespace(extract_text=lambda layout=False: "PAGE_B")]
            def __enter__(self): return self
            def __exit__(self, *args): return None
        with patch.dict("sys.modules", {"pdfplumber":SimpleNamespace(open=lambda path: FakePdf())}):
            pdf_text, _ = ops._extract_document(pdf_path, self.docs)
        self.assertIn("Страница 1\nPAGE_A", pdf_text)
        self.assertIn("Страница 2\nPAGE_B", pdf_text)

        class ScannedPdf:
            pages = [SimpleNamespace(extract_text=lambda layout=False: None),
                     SimpleNamespace(extract_text=lambda layout=False: "  ")]
            def __enter__(self): return self
            def __exit__(self, *args): return None
        with patch.dict("sys.modules", {"pdfplumber":SimpleNamespace(open=lambda path: ScannedPdf())}):
            with self.assertRaisesRegex(ValueError, "OCR не выполнялся"):
                ops._extract_document(pdf_path, self.docs)

    def test_document_extraction_is_complete_deduplicated_and_has_no_published_date(self):
        doc_id = "12345678-1234-5678-1234-567812345678"
        path = self.docs / f"{doc_id}.txt"
        text = "Строка 1\n" + "Важно.\n" * 2000
        path.write_text(text, encoding="utf-8")
        first = ops.prepare_document(doc_id, self.docs, self.db)
        second = ops.prepare_document(doc_id, self.docs, self.db)
        self.assertFalse(first["duplicate"])
        self.assertTrue(second["duplicate"])
        item = ops.get_inbox_item(first["inbox_id"], self.db)
        self.assertEqual(item["text"], text)
        self.assertIsNone(item["published_at"])
        self.assertEqual(item["metadata"]["text_chars"], len(text))
        self.assertEqual(len(ops.list_inbox("queued", 5, self.db)), 1)
        self.assertEqual(ops.update_inbox_state(item["id"], "reviewed", self.db)["state"], "reviewed")

    def test_json_document_is_preserved_and_empty_or_oversize_inputs_fail_visibly(self):
        doc_id = "12345678-1234-5678-1234-567812345679"
        path = self.docs / f"{doc_id}.json"
        path.write_text('{"title":"A","body":"Текст"}', encoding="utf-8")
        result = ops.prepare_document(doc_id, self.docs, self.db)
        self.assertIn('"body": "Текст"', ops.get_inbox_item(result["inbox_id"], self.db)["text"])
        empty_id = "12345678-1234-5678-1234-567812345680"
        (self.docs / f"{empty_id}.txt").write_text("  ", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "не найден извлекаемый текст"):
            ops.prepare_document(empty_id, self.docs, self.db)
        big_id = "12345678-1234-5678-1234-567812345681"
        (self.docs / f"{big_id}.txt").write_bytes(b"x" * (ops.MAX_DOCUMENT_BYTES + 1))
        with self.assertRaisesRegex(ValueError, "превышает лимит"):
            ops.prepare_document(big_id, self.docs, self.db)

    def test_document_parser_errors_include_reason_without_exposing_storage_path(self):
        parser_error = RuntimeError(f"Invalid PDF under {ops.DOCUMENTS_DIR}/source.pdf")
        with patch("backend.api.routes.operations.ops.prepare_document", side_effect=parser_error):
            with self.assertRaises(HTTPException) as caught:
                asyncio.run(operations_route.prepare_document("some-id", {}))
        self.assertEqual(caught.exception.status_code, 422)
        self.assertIn("RuntimeError", caught.exception.detail)
        self.assertIn("Invalid PDF", caught.exception.detail)
        self.assertNotIn(str(ops.DOCUMENTS_DIR), caught.exception.detail)

    def test_upload_is_bounded_and_accepts_markdown_without_trusting_filename(self):
        upload = ops.save_document_upload("../../notes.md", b"# Heading\nText", self.docs, title="Материалы штаба", db_path=self.db)
        self.assertEqual(upload["title"], "Материалы штаба")
        self.assertEqual(upload["status"], "uploaded")
        self.assertEqual((self.docs / f"{upload['id']}.md").read_text(), "# Heading\nText")
        self.assertEqual(stat.S_IMODE(os.stat(self.docs / f"{upload['id']}.md").st_mode), 0o600)
        prepared = ops.prepare_document(upload["id"], self.docs, self.db)
        self.assertEqual(ops.get_inbox_item(prepared["inbox_id"], self.db)["text"], "# Heading\nText")
        document_source = next(item for item in ops.list_sources(self.db) if item["id"] == prepared["source_id"])
        self.assertEqual(document_source["title"], "Материалы штаба")
        with self.assertRaises(ValueError): ops.save_document_upload("notes.exe", b"x", self.docs)
        with self.assertRaisesRegex(ValueError, "превышает лимит"):
            ops.save_document_upload("large.txt", b"x" * (ops.MAX_DOCUMENT_BYTES + 1), self.docs)

    def test_schedule_defaults_off_validates_sources_and_enqueues_one_daily_batch_per_day(self):
        default = ops.get_schedule(self.db)
        self.assertFalse(default["enabled"])
        self.assertEqual(default["timezone"], "Europe/Moscow")
        source = ops.add_vk_source("https://vk.com/dainima", self.db)
        ops._schedule_due(self.db)
        self.assertEqual(ops.list_jobs(20, self.db), [])
        with self.assertRaises(ValueError): ops.set_schedule(True, "08:00", "UTC", [source["id"]], self.db)
        ops.set_schedule(True, "00:00", "Europe/Moscow", [source["id"]], self.db)
        ops._schedule_due(self.db)
        first = ops.list_jobs(20, self.db)
        self.assertEqual(len(first), 1)
        self.assertEqual(first[0]["result"]["trigger"], "daily_schedule")
        ops._schedule_due(self.db)
        self.assertEqual(len(ops.list_jobs(20, self.db)), 1)
        ops.set_schedule(False, "00:00", "Europe/Moscow", [], self.db)
        self.assertFalse(ops.get_schedule(self.db)["enabled"])

    def test_startup_recovers_running_jobs_as_interrupted_without_erasing_history(self):
        source = ops.add_vk_source("https://vk.com/dainima", self.db)
        job = ops.create_vk_job(source["id"], 0, self.db)
        con = ops.connect(self.db)
        con.execute("UPDATE operations_jobs SET status='running',attempts=1 WHERE id=?", (job["id"],))
        con.commit(); con.close()
        self.assertEqual(ops.recover_interrupted_jobs(self.db), 1)
        recovered = ops.get_job(job["id"], self.db)
        self.assertEqual(recovered["status"], "interrupted")
        self.assertEqual(recovered["attempts"], 1)
        self.assertIn("stopped before", recovered["error"])
        retry = ops.retry_job(job["id"], self.db)
        self.assertEqual(retry["retry_of"], job["id"])
        self.assertEqual(ops.get_job(job["id"], self.db)["status"], "interrupted")

    def test_vk_collection_mock_records_source_in_inbox_and_posts_jsonl(self):
        source = ops.add_vk_source("https://vk.com/dainima", self.db)
        path = self.root / "posts" / "posts_dainima.jsonl"
        raw = [{"post_id":"101","post_url":"https://vk.com/dainima?w=wall-55_101","group_id":"55",
                "group_name":"Дайнима","group_domain":"dainima","published_at":"2026-10-17T12:00:00+03:00",
                "text_raw":"Подтвержденная новость", "attachments":{"photos":[],"links":[],"docs":[]}}]
        class FakeExtractor:
            def __init__(self, token): pass
            def resolve_group(self, url): return {"group_id":55,"group_domain":"dainima","group_name":"Дайнима"}
            def fetch_posts(self, limit, **kwargs): return iter(raw)
        # Avoid secrets and network: patch only the external VK extractor and corpus path.
        with patch("backend.admin.operations.STORAGE_DIR", self.root), \
             patch("backend.admin.secrets.get_secret", return_value="unit-test-token"):
            result = ops._collect_vk(source, 1, "job-test", self.db, FakeExtractor,
                                    lambda text: (text, [], []))
        self.assertEqual(result["fetched"], 1)
        self.assertEqual(result["parsed"], 1)
        self.assertEqual(result["inbox_added"], 1)
        saved = json.loads(path.read_text(encoding="utf-8").splitlines()[0])
        self.assertEqual(saved["text_raw"], "Подтвержденная новость")
        items = ops.list_inbox("queued", 10, self.db)
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["external_id"], "101")
        self.assertEqual(items[0]["published_at"], "2026-10-17T12:00:00+03:00")
        self.assertEqual(items[0]["content_sha256"], hashlib.sha256(saved["text_raw"].encode()).hexdigest())

    def test_vk_collection_recovers_a_crash_after_jsonl_append_before_inbox_receipt(self):
        source = ops.add_vk_source("https://vk.com/dainima", self.db)
        source_path = self.root / "posts" / "posts_dainima.jsonl"
        source_path.parent.mkdir(parents=True)
        raw = {"post_id":"101","post_url":"https://vk.com/dainima?w=wall-55_101","group_id":"55",
               "group_name":"Дайнима","group_domain":"dainima","published_at":"2026-10-17T12:00:00+03:00",
               "text_raw":"Уже сохранено в jsonl", "text_clean":"Уже сохранено", "attachments":{"photos":[],"links":[],"docs":[]}}
        source_path.write_text(json.dumps(raw, ensure_ascii=False)+"\n", encoding="utf-8")
        class FakeExtractor:
            def __init__(self, token): pass
            def resolve_group(self, url): return {"group_id":55,"group_domain":"dainima","group_name":"Дайнима"}
            def fetch_posts(self, limit, **kwargs): return iter([{k:v for k,v in raw.items() if k not in {"text_clean"}}])
        with patch("backend.admin.operations.STORAGE_DIR", self.root), \
             patch("backend.admin.secrets.get_secret", return_value="unit-test-token"):
            result = ops._collect_vk(source, 1, "recovery-test", self.db, FakeExtractor,
                                    lambda text: (text, [], []))
        self.assertEqual(result["fetched"], 0)
        self.assertEqual(result["parsed"], 1)
        self.assertEqual(result["inbox_added"], 1)
        self.assertEqual(len(ops.list_inbox("queued", 5, self.db)), 1)

    def test_vk_collection_separates_a_crash_truncated_jsonl_tail(self):
        source = ops.add_vk_source("https://vk.com/dainima", self.db)
        source_path = self.root / "posts" / "posts_dainima.jsonl"
        source_path.parent.mkdir(parents=True)
        source_path.write_bytes(b'{"partial":')
        raw = {"post_id":"102","post_url":"https://vk.com/dainima?w=wall-55_102","group_id":"55",
               "group_name":"Дайнима","published_at":"2026-10-17T12:00:00+03:00","text_raw":"Новая запись"}
        class FakeExtractor:
            def __init__(self, token): pass
            def resolve_group(self, url): return {"group_id":55,"group_domain":"dainima","group_name":"Дайнима"}
            def fetch_posts(self, limit, **kwargs): return iter([raw])
        with patch("backend.admin.operations.STORAGE_DIR", self.root), \
             patch("backend.admin.secrets.get_secret", return_value="unit-test-token"):
            result = ops._collect_vk(source, 1, "tail-recovery", self.db, FakeExtractor,
                                     lambda text: (text, [], []))
        rows = source_path.read_text(encoding="utf-8").splitlines()
        self.assertEqual(len(rows), 2)
        self.assertEqual(json.loads(rows[1])["post_id"], "102")
        self.assertEqual(result["fetched"], 1)

    def test_job_runner_persists_success_and_failure_results(self):
        source = ops.add_vk_source("https://vk.com/dainima", self.db)
        good = ops.create_vk_job(source["id"], 4, self.db)
        with patch("backend.admin.operations._collect_vk", return_value={"fetched":3,"parsed":3,"inbox_added":3}):
            ops._run_job(good["id"], self.db)
        done = ops.get_job(good["id"], self.db)
        self.assertEqual(done["status"], "succeeded")
        self.assertEqual(done["attempts"], 1)
        self.assertEqual(done["result"]["fetched"], 3)
        # A duplicate queue observation cannot claim or rerun a completed job.
        ops._run_job(good["id"], self.db)
        self.assertEqual(ops.get_job(good["id"], self.db)["attempts"], 1)
        bad = ops.create_vk_job(source["id"], 0, self.db)
        with patch("backend.admin.operations._collect_vk", side_effect=RuntimeError("mock fetch failure")):
            ops._run_job(bad["id"], self.db)
        failed = ops.get_job(bad["id"], self.db)
        self.assertEqual(failed["status"], "failed")
        self.assertIn("mock fetch failure", failed["error"])

    def test_job_error_redacts_configured_vk_token(self):
        source = ops.add_vk_source("https://vk.com/dainima", self.db)
        job = ops.create_vk_job(source["id"], 0, self.db)
        with patch("backend.admin.operations._collect_vk", side_effect=RuntimeError("bad token=super-secret-value")), \
             patch("backend.admin.secrets.get_secret", return_value="super-secret-value"):
            ops._run_job(job["id"], self.db)
        error = ops.get_job(job["id"], self.db)["error"]
        self.assertNotIn("super-secret-value", error)
        self.assertIn("[REDACTED]", error)


if __name__ == "__main__":
    unittest.main()
